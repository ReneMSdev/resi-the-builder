import os
import shutil
import subprocess

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

PAGE_MARGIN_INCHES = 0.75
USABLE_WIDTH_INCHES = 8.5 - (PAGE_MARGIN_INCHES * 2)  # US Letter width minus margins

# Hardcoded to the default macOS Homebrew cask install location. If this ever runs on a
# different machine or OS, this path needs to change.
SOFFICE_PATH = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def _add_entry_header(doc, entry: dict):
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(USABLE_WIDTH_INCHES), WD_TAB_ALIGNMENT.RIGHT
    )
    title = entry.get("title", "")
    organization = entry.get("organization", "")
    header_text = " — ".join(part for part in (title, organization) if part)
    run = p.add_run(header_text)
    run.bold = True

    dates = entry.get("dates", "")
    if dates:
        p.add_run(f"\t{dates}")

    location = entry.get("location", "")
    if location:
        loc_p = doc.add_paragraph()
        loc_run = loc_p.add_run(location)
        loc_run.italic = True
        loc_run.font.size = Pt(9.5)


def _add_section_heading(doc, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(12)


def render_resume_docx(resume: dict, output_path: str) -> str:
    """Builds a .docx file at output_path from a Resume dict. Returns output_path."""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Inches(PAGE_MARGIN_INCHES)
        section.bottom_margin = Inches(PAGE_MARGIN_INCHES)
        section.left_margin = Inches(PAGE_MARGIN_INCHES)
        section.right_margin = Inches(PAGE_MARGIN_INCHES)

    meta = resume.get("meta", {})

    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(meta.get("name", ""))
    name_run.bold = True
    name_run.font.size = Pt(20)

    contact_parts = [p for p in (meta.get("email"), meta.get("phone")) if p]
    for link in meta.get("links", []):
        label = link.get("label", "")
        url = link.get("url", "")
        contact_parts.append(f"{label}: {url}" if label else url)

    if contact_parts:
        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_run = contact_p.add_run(" | ".join(contact_parts))
        contact_run.font.size = Pt(9.5)

    summary = resume.get("summary") or {}
    summary_text = (summary.get("text") or "").strip()
    if summary_text:
        summary_p = doc.add_paragraph(summary_text)
        summary_p.paragraph_format.space_before = Pt(8)

    for section in resume.get("sections", []):
        entries = section.get("entries") or []
        groups = section.get("groups") or []
        if not entries and not groups:
            continue

        _add_section_heading(doc, section.get("title", ""))
        sec_type = section.get("type", "")

        if sec_type == "skills":
            for group in groups:
                p = doc.add_paragraph()
                label_run = p.add_run(f"{group.get('label', '')}: ")
                label_run.bold = True
                p.add_run(", ".join(group.get("items", [])))

        elif sec_type in ("education", "certifications"):
            for entry in entries:
                parts = [
                    entry.get("title", ""),
                    entry.get("organization", ""),
                    entry.get("dates", ""),
                ]
                doc.add_paragraph(" — ".join(part for part in parts if part))

        else:
            # "experience", "projects", and any other entry-shaped section type
            for entry in entries:
                _add_entry_header(doc, entry)
                for bullet in entry.get("bullets", []):
                    doc.add_paragraph(bullet.get("text", ""), style="List Bullet")

    doc.save(output_path)
    return output_path


def render_cover_letter_docx(cover_letter: dict, output_path: str) -> str:
    """Builds a business-letter-style .docx file at output_path from a CoverLetter dict.
    Returns output_path."""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Inches(PAGE_MARGIN_INCHES)
        section.bottom_margin = Inches(PAGE_MARGIN_INCHES)
        section.left_margin = Inches(PAGE_MARGIN_INCHES)
        section.right_margin = Inches(PAGE_MARGIN_INCHES)

    meta = cover_letter.get("meta", {})

    date = (meta.get("date") or "").strip()
    if date:
        date_p = doc.add_paragraph()
        date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        date_p.add_run(date)

    role = meta.get("role", "")
    company = meta.get("company", "")
    subject = " at ".join(part for part in (role, company) if part)
    if subject:
        subject_p = doc.add_paragraph()
        subject_p.paragraph_format.space_before = Pt(8)
        subject_run = subject_p.add_run(f"Re: {subject}")
        subject_run.bold = True

    salutation_p = doc.add_paragraph("Dear Hiring Manager,")
    salutation_p.paragraph_format.space_before = Pt(8)

    for paragraph in cover_letter.get("paragraphs", []):
        text = (paragraph.get("text") or "").strip()
        if not text:
            continue
        p = doc.add_paragraph(text)
        p.paragraph_format.space_before = Pt(8)

    sign_off_p = doc.add_paragraph("Sincerely,")
    sign_off_p.paragraph_format.space_before = Pt(8)

    name_p = doc.add_paragraph()
    name_run = name_p.add_run(meta.get("name", ""))
    name_run.bold = True

    contact_parts = [p for p in (meta.get("email"), meta.get("phone")) if p]
    if contact_parts:
        contact_p = doc.add_paragraph()
        contact_run = contact_p.add_run(" | ".join(contact_parts))
        contact_run.font.size = Pt(9.5)

    doc.save(output_path)
    return output_path


def convert_docx_to_pdf(docx_path: str, output_dir: str) -> str:
    """Converts a .docx file to .pdf using headless LibreOffice.
    Returns the path to the resulting .pdf file."""
    if not shutil.which(SOFFICE_PATH) and not os.path.exists(SOFFICE_PATH):
        raise RuntimeError(
            "LibreOffice not found at expected path. Install with "
            "'brew install --cask libreoffice' or update SOFFICE_PATH."
        )

    result = subprocess.run(
        [
            SOFFICE_PATH,
            "--headless",
            "--convert-to", "pdf",
            "--outdir", output_dir,
            docx_path,
        ],
        capture_output=True,
        text=True,
        timeout=30,
    )

    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr}")

    pdf_path = os.path.join(
        output_dir,
        os.path.splitext(os.path.basename(docx_path))[0] + ".pdf"
    )
    if not os.path.exists(pdf_path):
        raise RuntimeError("LibreOffice reported success but no PDF file was produced.")

    return pdf_path
